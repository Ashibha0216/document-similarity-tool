from flask import Flask, render_template, request
import os
import re
import pymupdf
import webbrowser
from threading import Timer
from docx import Document
from pptx import Presentation
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

app = Flask(__name__)

UPLOAD_FOLDER = "uploads"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

def open_browser():
    webbrowser.open("http://127.0.0.1:5000")

def extract_text(filepath):
    extension = os.path.splitext(filepath)[1].lower()

    if extension == ".pdf":
        text = ""
        with pymupdf.open(filepath) as pdf:
            for page in pdf:
                text += page.get_text()
        return text

    elif extension == ".docx":
        document = Document(filepath)
        return "\n".join(
            paragraph.text for paragraph in document.paragraphs
        )

    elif extension == ".pptx":
        presentation = Presentation(filepath)
        text = ""

        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"

        return text

    return ""


def clean_text(text):
    text = text.lower()
    text = re.sub(r"[^a-zA-Z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def calculate_similarity(text1, text2):

    text1 = clean_text(text1)
    text2 = clean_text(text2)

    if not text1 or not text2:
        return 0

    vectorizer = TfidfVectorizer(stop_words="english")

    matrix = vectorizer.fit_transform([text1, text2])

    similarity = cosine_similarity(matrix[0:1], matrix[1:2])[0][0]

    return similarity * 100


@app.route("/", methods=["GET", "POST"])
def home():

    if request.method == "POST":

        file1 = request.files.get("file1")
        file2 = request.files.get("file2")

        if not file1 or not file2:
            return render_template(
                "index.html",
                error="Please select both files."
            )

        file1_path = os.path.join(UPLOAD_FOLDER, file1.filename)
        file2_path = os.path.join(UPLOAD_FOLDER, file2.filename)

        file1.save(file1_path)
        file2.save(file2_path)

        try:
            text1 = extract_text(file1_path)
            text2 = extract_text(file2_path)

            similarity = calculate_similarity(text1, text2)

            return render_template(
                "index.html",
                result=True,
                similarity=f"{similarity:.2f}",
                file1_name=file1.filename,
                file2_name=file2.filename
            )

        except Exception as e:
            return render_template(
                "index.html",
                error=f"Error: {str(e)}"
            )

    return render_template("index.html")


if __name__ == "__main__":
    Timer(1.5, open_browser).start()
    app.run(debug=False, use_reloader=False)